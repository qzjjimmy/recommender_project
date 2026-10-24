"""把课程报告中的项目三旧状态更新为当前已实现状态。

输入报告保留不动，输出一份新的 DOCX，供重新渲染和最终提交。
"""
from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "docs"
    / "reports"
    / "《行业工程实践》专升本学生实训报告（项目一至三已完成）-表格字体统一版.docx"
)
OUTPUT = (
    ROOT
    / "docs"
    / "reports"
    / "《行业工程实践》专升本学生实训报告（项目一至三已完成）-项目三验收补齐版.docx"
)


def replace_paragraph(paragraph: Paragraph, text: str) -> None:
    """替换整段文字，并沿用第一个 run 的字符格式。"""
    template_rpr = None
    for run in paragraph.runs:
        if run.text.strip():
            rpr = run._r.rPr
            template_rpr = copy.deepcopy(rpr) if rpr is not None else None
            break
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if template_rpr is not None:
        run._r.insert(0, template_rpr)


def replace_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    template_rpr = None
    if paragraph.runs:
        rpr = paragraph.runs[0]._r.rPr
        template_rpr = copy.deepcopy(rpr) if rpr is not None else None
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if template_rpr is not None:
        run._r.insert(0, template_rpr)


def replace_starting(doc: Document, prefix: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            replace_paragraph(paragraph, text)
            return
    raise ValueError(f"未找到待更新段落：{prefix}")


def find_table(doc: Document, first_header: str):
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == first_header:
            return table
    raise ValueError(f"未找到表格：{first_header}")


def update_fixed_tables(doc: Document) -> None:
    interfaces = find_table(doc, "接口/模块")
    interface_rows = [
        (
            "多模态助手→推荐服务",
            "trace_id、multipart image、可选 query；图片≤8MB",
            "trace_id、scene、objects、colors、category_hints、suggested_query",
            "浏览器→Flask→桥接/vision；X-Trace-ID 双向回传",
            "格式错误400；过大413；桥接或模型不可用503，保留文字与规则推荐",
        ),
        (
            "推荐服务→排序编排",
            "trace_id、user_id、session_id、timestamp、source、q、top_n",
            "intent、status、constraints、多路来源、EvidencePack、Top-K",
            "InteractionEvent→SessionContext→Orchestrator→SmartSearch",
            "缺槽位追问；冲突澄清；空结果提示；硬约束先于打分",
        ),
        (
            "解释模块",
            "商品属性、命中约束、召回来源、置信度≥0.6的视觉标签",
            "唯一 trace、逐商品证据、reason 与 reason_detail",
            "EvidencePack→确定性解释器",
            "无证据明确返回信息不足，不补写价格、品牌或属性",
        ),
        (
            "反馈模块",
            "username、session_id、item_id、action(like/dislike/browse)",
            "state、stats、hint、当前会话上下文",
            "按钮/手势/浏览→/feedback 或 /api/browse→SessionContext",
            "点赞长期保存；点踩与浏览仅当前会话；互斥、可撤销、可清空",
        ),
    ]
    for row, values in zip(interfaces.rows[1:5], interface_rows):
        for cell, value in zip(row.cells, values):
            replace_cell(cell, value)

    e2e = find_table(doc, "用例")
    for row in e2e.rows[1:]:
        if row.cells[0].text.strip() == "E01":
            replace_cell(
                row.cells[4],
                "通过；文字、语音、视觉均进入统一编排；同输入固定种子可复现",
            )
            replace_cell(
                row.cells[5],
                "outputs/stage03_demo.json；outputs/stage06_acceptance.json",
            )
        elif row.cells[0].text.strip() == "E04":
            replace_cell(
                row.cells[4],
                "通过；桥接故障返回503与text-only降级，文字搜索保持200",
            )
            replace_cell(
                row.cells[5],
                "outputs/stage06_acceptance.json；tests/test_multimodal_trace.py",
            )

    delivery = find_table(doc, "交付物")
    for row in delivery.rows[1:]:
        label = row.cells[0].text.strip()
        if label == "各阶段输出、日志与必要截图":
            replace_cell(
                row.cells[2],
                "outputs/stage01–stage06 结构化产物与 stage22 视觉证据",
            )
            replace_cell(
                row.cells[3],
                "契约、26条上下文样例、三类trace、证据表、反馈A/B、端到端验收",
            )
        elif label == "架构图与接口说明":
            replace_cell(row.cells[2], "docs/architecture.md；第5章")
            replace_cell(
                row.cells[3],
                "上下文、容器、时序、统一协议、会话和失败出口",
            )
        elif label == "部署/运行说明与风险登记":
            replace_cell(row.cells[2], "docs/deployment_and_risks.md；README.md")
            replace_cell(
                row.cells[3],
                "启动、重跑、模型下载、反馈隔离、隐私和提交包清单",
            )

    personal = None
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == "序号":
            personal = table
            break
    if personal is None:
        raise ValueError("未找到个人任务表")
    for row in personal.rows[1:]:
        number = row.cells[0].text.strip()
        if number == "3":
            replace_cell(row.cells[1], "项目三接口、会话与编排")
            replace_cell(
                row.cells[2],
                "trace协议、六类意图、SessionContext、EvidencePack与Orchestrator",
            )
            replace_cell(
                row.cells[3],
                "src/protocol.py、session_context.py、evidence.py、orchestrator.py",
            )
        elif number == "4":
            replace_cell(row.cells[1], "测试、部署与提交包")
            replace_cell(
                row.cells[2],
                "六阶段重跑脚本、会话隔离A/B、三类输入、故障降级与打包审计",
            )
            replace_cell(
                row.cells[3],
                "scripts/01–06；tests/；docs/；dist/",
            )

    evidence = find_table(doc, "证据类型")
    evidence_rows = [
        (
            "Git工作区",
            "main / 2026-10-24",
            "当前工作树；模型与敏感文件已移出Git跟踪",
            "版本追溯",
            "不虚构提交哈希；以文件和阶段产物复核",
        ),
        (
            "六阶段验收",
            "2026-10-24",
            "scripts/01_contract_test.py 至 06_e2e_acceptance.py",
            "契约/上下文/编排/解释/反馈/联调",
            "按顺序重跑并落盘JSON、JSONL、CSV与Markdown",
        ),
        (
            "自动化测试",
            "2026-10-24",
            "tests/；multimodal_ai_project1/tests/",
            "协议、会话、证据、编排、网页与桥接trace",
            "根项目26项与桥接项目5项测试通过",
        ),
        (
            "提交包",
            "2026-10-24",
            "dist/project3_submission.zip",
            "最终交付",
            "显式清单排除模型、数据库、密钥、缓存和日志",
        ),
    ]
    for row, values in zip(evidence.rows[1:5], evidence_rows):
        for cell, value in zip(row.cells, values):
            replace_cell(cell, value)


def main() -> int:
    doc = Document(SOURCE)
    updates = [
        (
            "项目三达到的主要结果包括：",
            "项目三达到的主要结果包括：两项目通过HTTP和X-Trace-ID协作；六个阶段均有可重跑入口；"
            "26条文本、视觉和手势组合样例通过；Orchestrator保留多路召回来源并把硬约束放在打分前；"
            "EvidencePack为推荐提供逐商品事实；点赞长期保存，点踩与浏览按会话隔离；"
            "桥接故障明确降级而文字推荐继续可用。",
        ),
        (
            "版本管理现状：",
            "版本管理现状：交付目录是Git仓库，当前工作树保留本次新增实现。根目录与项目二均通过"
            ".gitignore排除模型、数据库、密钥、缓存和运行日志；MediaPipe的.task文件已移出Git跟踪，"
            "但仍保留在本机供演示使用。提交包由显式清单生成，不依赖通配复制。",
        ),
        (
            "与项目三的衔接：",
            "与项目三的衔接：项目三通过本地桥接使用SenseVoiceSmall、MediaPipe、Qwen2.5-VL、"
            "Ollama与edge-tts。语音文本、视觉线索和手势动作统一转换为InteractionEvent；"
            "同一会话内点踩和浏览立即影响下一次排序，点赞作为长期偏好保存。两项目以HTTP解耦，"
            "既不读取彼此内部变量，也避免同名config包冲突。",
        ),
        (
            "场景三是结果反馈。",
            "场景三是结果反馈。用户对商品点赞、点踩、继续浏览，或用手势表达喜欢与不喜欢。"
            "点赞写入用户级SQLite作为长期偏好；点踩和浏览只保存在当前SessionContext，"
            "新会话不会继承旧会话的负反馈。系统不识别个人身份，视觉请求不保存原图。",
        ),
        (
            "语音输入先由SenseVoiceSmall转写成UTF-8文本",
            "图片输入由视觉模型输出可见物体、颜色、合法类目和建议搜索词。视觉选品只使用图片中"
            "可核验的物品与类目；预算和价格由用户在可编辑搜索框中单独输入，不把模型猜测的价格"
            "当成商品事实。颜色作为证据展示，默认不加入查询，避免高频颜色稀释物体线索。",
        ),
        (
            "项目三没有新增复杂深度排序模型",
            "项目三没有新增复杂深度排序模型，而是复用项目一的可解释流程。SessionContext先融合"
            "本轮或追问的条件，Orchestrator再调用SmartSearch执行品类、价格与关键词硬过滤。"
            "排序仍为“0.50×相似度 + 0.25×热度 + 0.25×反馈”，每条结果保留召回来源与EvidencePack；"
            "低置信度视觉标签不进入偏好或解释。",
        ),
        (
            "点赞、点踩和手势反馈都调用同一个/feedback接口。",
            "点赞、点踩和手势反馈都调用同一个/feedback接口。点赞写入用户级SQLite并可长期复用；"
            "点踩与继续浏览只进入当前session_id，保证负反馈不会污染其他会话。重复同一动作表示取消，"
            "喜欢与不喜欢互斥。会话可通过/api/session/context查看，并支持撤销与清理。",
        ),
        (
            "项目未实现跨服务trace_id。",
            "跨服务trace_id已经实现。Flask在API入口生成或接收X-Trace-ID，调用桥接时继续传递，"
            "桥接响应头和JSON会返回同一trace；推荐请求、意图、召回来源、证据和错误响应都使用该字段。"
            "六阶段验收在outputs/stage01_contract_report.json和stage06_acceptance.json中验证了传递结果。",
        ),
        (
            "系统处理的是教学演示与合成交易数据",
            "系统处理的是教学演示与合成交易数据，仍按真实数据风险设计控制。视觉图片只在内存和"
            "本地HTTP请求中使用，桥接端重新编码去除EXIF，不保存原图。账号密码使用PBKDF2加盐哈希；"
            "用户级库只持久化点赞，点踩与浏览按会话隔离；会话密钥、模型密钥和本地数据库不进入提交包。",
        ),
        (
            "可验证的改进包括：",
            "可验证的改进包括：跨conda环境的HTTP桥接与双向trace；六类意图和带证据来源的槽位；"
            "可撤销的SessionContext；保留多路来源的Orchestrator；逐商品EvidencePack；"
            "点赞长期化与负反馈会话隔离；固定口径反馈A/B；三类输入和桥接故障降级；"
            "以及不含模型和敏感数据的显式提交包。",
        ),
        (
            "本人负责推荐系统、项目二核心模块和项目三整合",
            "本人负责推荐系统、项目二核心模块和项目三整合：推荐侧完成清洗、Item2Vec、召回、"
            "LR/LightGBM、Flask服务与反馈首页；多模态侧完成ASR、LLM/RAG、视觉、TTS与界面；"
            "项目三新增统一协议、SessionContext、Orchestrator、EvidencePack、会话反馈隔离、"
            "六阶段验收脚本和提交包审计。",
        ),
        (
            "核心实现三是项目一的视觉查询。",
            "核心实现三是图片到推荐的安全衔接。visual_intent只保留画面中可核验的物体、颜色和"
            "真实商品类目，默认使用主类目与最多三个物体名形成建议查询；预算由用户在可编辑搜索框"
            "单独输入，视觉结果不携带虚假价格。前端vision.js负责拍照、选图、浏览器缩放、取消、"
            "预览与确认，原图仅在本地请求中使用。",
        ),
        (
            "本目录尚未初始化为Git仓库",
            "交付目录现为Git工作区，本次验收不虚构提交哈希。证据以可以重跑的01至06脚本及其"
            "JSON、JSONL、CSV、Markdown产物为准；真实视觉链路仍由stage22证据补充，"
            "确定性阶段6验收使用桥接替身覆盖ASR、手势、视觉与故障降级。",
        ),
        (
            "接下来我会继续补OCR和连续帧稳定性",
            "接下来我会继续补OCR和连续帧稳定性，让视觉读取真实商品标签；把点赞数据回流到精排训练；"
            "增加跨服务更细的延迟分段；扩展多用户并发和长时稳定性测试；并准备本地离线TTS。"
            "当前版本已经把“能演示”推进到“有协议、有证据、能降级、能按阶段重跑”。",
        ),
    ]
    for prefix, text in updates:
        replace_starting(doc, prefix, text)
    update_fixed_tables(doc)
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("toc") and paragraph.text.strip().startswith(
            "8、实训总结与学习体会"
        ):
            replace_paragraph(paragraph, "8、实训总结与学习体会\t41")
            break
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    doc.core_properties.comments = (
        "项目一、项目二与项目三已完成；本版已同步统一协议、会话隔离、"
        "EvidencePack、六阶段验收和提交包审计结果。"
    )
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
